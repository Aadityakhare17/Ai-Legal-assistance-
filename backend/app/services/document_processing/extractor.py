"""
Document text extraction service.
Supports PDF (via PyMuPDF), DOCX (via python-docx), and TXT files.
"""
import os
import re
from pathlib import Path
from typing import Dict, List, Optional, Tuple
from loguru import logger


def extract_text_from_file(file_path: str) -> Tuple[str, int, int]:
    """
    Extract text from a file. Returns (text, page_count, word_count).
    """
    path = Path(file_path)
    ext = path.suffix.lower()

    if ext == ".pdf":
        return _extract_from_pdf(file_path)
    elif ext in (".docx", ".doc"):
        return _extract_from_docx(file_path)
    elif ext == ".txt":
        return _extract_from_txt(file_path)
    else:
        raise ValueError(f"Unsupported file type: {ext}")


def _extract_from_pdf(file_path: str) -> Tuple[str, int, int]:
    """Extract text from PDF using PyMuPDF."""
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(file_path)
        text_parts = []
        page_count = len(doc)

        for page_num, page in enumerate(doc):
            page_text = page.get_text("text")
            if page_text.strip():
                text_parts.append(f"[Page {page_num + 1}]\n{page_text}")

        full_text = "\n\n".join(text_parts)

        # If PDF is scanned (no extractable text), indicate it
        if not full_text.strip():
            full_text = "[SCANNED_PDF] This document appears to be a scanned image. Text extraction was limited."

        word_count = len(full_text.split())
        doc.close()
        return full_text, page_count, word_count
    except Exception as e:
        logger.error(f"PDF extraction error: {e}")
        raise


def _extract_from_docx(file_path: str) -> Tuple[str, int, int]:
    """Extract text from DOCX using python-docx."""
    try:
        from docx import Document
        doc = Document(file_path)
        paragraphs = []

        for para in doc.paragraphs:
            if para.text.strip():
                paragraphs.append(para.text)

        # Also extract from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text)

        full_text = "\n\n".join(paragraphs)
        word_count = len(full_text.split())
        return full_text, 1, word_count  # DOCX doesn't have fixed pages
    except Exception as e:
        logger.error(f"DOCX extraction error: {e}")
        raise


def _extract_from_txt(file_path: str) -> Tuple[str, int, int]:
    """Extract text from TXT file."""
    try:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            text = f.read()
        word_count = len(text.split())
        return text, 1, word_count
    except Exception as e:
        logger.error(f"TXT extraction error: {e}")
        raise


def chunk_text(text: str, chunk_size: int = 1000, overlap: int = 200) -> List[Dict]:
    """
    Split text into overlapping chunks for RAG embedding.
    Returns list of {content, chunk_index, page_number, section_heading}
    """
    chunks = []
    lines = text.split("\n")
    current_chunk = []
    current_length = 0
    chunk_index = 0
    current_page = None
    current_heading = None

    for line in lines:
        # Track page markers from PDF extraction
        page_match = re.match(r"\[Page (\d+)\]", line)
        if page_match:
            current_page = int(page_match.group(1))
            continue

        # Detect section headings (ALL CAPS or numbered sections)
        if re.match(r"^[A-Z][A-Z\s]{5,}$", line.strip()) or re.match(r"^\d+\.\s+[A-Z]", line.strip()):
            current_heading = line.strip()

        words = line.split()
        current_chunk.append(line)
        current_length += len(words)

        if current_length >= chunk_size:
            chunk_text_content = "\n".join(current_chunk)
            chunks.append({
                "content": chunk_text_content,
                "chunk_index": chunk_index,
                "page_number": current_page,
                "section_heading": current_heading,
            })
            chunk_index += 1
            # Keep overlap
            overlap_lines = current_chunk[-max(1, len(current_chunk) // 5):]
            current_chunk = overlap_lines
            current_length = sum(len(l.split()) for l in overlap_lines)

    # Add remaining
    if current_chunk:
        chunks.append({
            "content": "\n".join(current_chunk),
            "chunk_index": chunk_index,
            "page_number": current_page,
            "section_heading": current_heading,
        })

    return chunks


def validate_file(filename: str, file_size: int, allowed_extensions: List[str], max_size_mb: int) -> Optional[str]:
    """Validate file before processing. Returns error message or None."""
    ext = Path(filename).suffix.lower()
    if ext not in allowed_extensions:
        return f"File type '{ext}' is not supported. Allowed types: {', '.join(allowed_extensions)}"
    if file_size > max_size_mb * 1024 * 1024:
        return f"File size exceeds the maximum allowed size of {max_size_mb}MB"
    return None
